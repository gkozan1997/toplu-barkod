"""
Toplu Barkod & 10x10cm Termal Kargo Etiket Motoru
Excel/CSV okuma, dinamik sütun tespiti ve ReportLab 100mm x 100mm PDF üretim motoru.
"""

import os
import io
import re
import csv
from datetime import datetime
import openpyxl
import xlrd
from reportlab.lib.pagesizes import mm
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether, PageBreak, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.graphics.barcode import createBarcodeDrawing
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FONTS_INITIALIZED = False

def init_fonts():
    global FONTS_INITIALIZED
    if FONTS_INITIALIZED:
        return
    local_fonts_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fonts')
    font_defs = [
        ('LabelFont', os.path.join(local_fonts_dir, 'arial.ttf')),
        ('LabelFont-Bold', os.path.join(local_fonts_dir, 'arialbd.ttf')),
        ('LabelSegoe', r'C:\Windows\Fonts\segoeui.ttf'),
        ('LabelSegoe-Bold', r'C:\Windows\Fonts\segoeuib.ttf'),
        ('LabelFont', r'C:\Windows\Fonts\arial.ttf'),
        ('LabelFont-Bold', r'C:\Windows\Fonts\arialbd.ttf'),
    ]
    for name, path in font_defs:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
            except Exception as e:
                print(f"Font yuklenemedi {name}: {e}")
    FONTS_INITIALIZED = True


def tr_lower(text):
    if not text:
        return ""
    return str(text).replace('İ', 'i').replace('I', 'ı').replace('Ğ', 'ğ').replace('Ü', 'ü').replace('Ş', 'ş').replace('Ö', 'ö').replace('Ç', 'ç').lower()


def clean_barcode(val):
    if val is None:
        return ""
    s = str(val).strip()
    if not s or s.lower() in ('none', 'nan', 'null', '-'):
        return ""
    if '.' in s or 'e' in s.lower():
        try:
            f = float(s)
            if f.is_integer():
                return str(int(f))
        except Exception:
            pass
    if s.endswith('.0'):
        s = s[:-2]
    return s.strip()


def clean_quantity(val):
    if val is None:
        return 1
    try:
        if isinstance(val, (int, float)):
            q = int(val)
            return q if q > 0 else 1
        s = str(val).strip()
        if not s:
            return 1
        q = int(float(s.replace(',', '.')))
        return q if q > 0 else 1
    except Exception:
        return 1


def clean_text(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ('none', 'nan', 'null'):
        return ""
    return re.sub(r'\s+', ' ', s)


# Sütun adı algılama anahtar kelimeleri
CUSTOMER_KEYWORDS = ['teslimat adı', 'teslimat adi', 'üye adı soyadı', 'uye adi soyadi', 'sevk - müşteri', 'sevk - musteri', 'alıcı', 'alici', 'müşteri', 'musteri', 'ad soyad', 'isim', 'recipient', 'customer', 'fatura adı', 'fatura adi']
ADDRESS_KEYWORDS = ['teslimat adresi', 'sevk - adres', 'adres', 'fatura adresi', 'adres satırı']
DISTRICT_KEYWORDS = ['teslimat ilçe adı', 'teslimat ilce adi', 'teslimat ilçe', 'teslimat ilce', 'sevk - ilçe', 'ilçe', 'ilce', 'district', 'fatura bölge adı']
CITY_KEYWORDS = ['teslimat şehir adı', 'teslimat sehir adi', 'teslimat şehir', 'teslimat sehir', 'sevk - il', 'şehir', 'sehir', 'il', 'city', 'fatura şehir adı']
PHONE_KEYWORDS = ['teslimat telefonu', 'müşteri telefon', 'musteri telefon', 'sevk - telefon', 'telefon', 'tel', 'gsm', 'fatura telefonu']
BARCODE_KEYWORDS = ['sipariş no', 'siparis no', 'kargo barkod', 'paket no', 'kargo takip no', 'barkod', 'barcode', 'gtin (barkod)', 'gtin', 'ean', 'ean13', 'barkod no']
NAME_KEYWORDS = ['ürün', 'ürün adı', 'urun', 'urun adi', 'product name', 'title', 'başlık', 'ürün ismi', 'urun ismi', 'ürün tanımı']
SKU_KEYWORDS = ['stok kodu', 'sku', 'ürün kodu', 'urun kodu', 'model kodu', 'stok']
QTY_KEYWORDS = ['miktar', 'adet', 'miktar (adet)', 'sipariş adedi', 'qty', 'quantity']


def match_column(headers, keyword_list):
    normalized_headers = [tr_lower(h).strip() if h is not None else "" for h in headers]
    for kw in keyword_list:
        for idx, h in enumerate(normalized_headers):
            if h == kw:
                return idx
    for kw in keyword_list:
        for idx, h in enumerate(normalized_headers):
            if kw in h:
                return idx
    return None


def read_excel_or_csv(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    raw_rows = []

    if ext == '.csv':
        for enc in ['utf-8-sig', 'utf-8', 'windows-1254', 'iso-8859-9', 'latin-1']:
            try:
                with open(file_path, 'r', encoding=enc) as f:
                    sample = f.read(2048)
                    f.seek(0)
                    dialect = csv.Sniffer().sniff(sample) if sample else 'excel'
                    reader = csv.reader(f, dialect=dialect)
                    raw_rows = [row for row in reader if any(cell.strip() for cell in row)]
                if raw_rows:
                    break
            except Exception:
                continue

    elif ext == '.xlsx':
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                raw_rows.append(list(row))
        wb.close()

    elif ext == '.xls':
        wb = xlrd.open_workbook(file_path)
        sheet = wb.sheet_by_index(0)
        for r in range(sheet.nrows):
            row_vals = sheet.row_values(r)
            if any(cell is not None and str(cell).strip() != "" for cell in row_vals):
                raw_rows.append(row_vals)

    if not raw_rows:
        return [], [], {}

    header_row_idx = 0
    best_match_count = 0

    for idx, row in enumerate(raw_rows[:10]):
        row_str = " ".join([tr_lower(c) for c in row if c is not None])
        matches = 0
        if any(kw in row_str for kw in BARCODE_KEYWORDS): matches += 2
        if any(kw in row_str for kw in NAME_KEYWORDS): matches += 2
        if any(kw in row_str for kw in CUSTOMER_KEYWORDS): matches += 2
        if any(kw in row_str for kw in ADDRESS_KEYWORDS): matches += 1
        if any(kw in row_str for kw in QTY_KEYWORDS): matches += 1
        if matches > best_match_count:
            best_match_count = matches
            header_row_idx = idx

    headers = [str(c).strip() if c is not None else f"Sütun_{i+1}" for i, c in enumerate(raw_rows[header_row_idx])]
    data_rows = raw_rows[header_row_idx + 1:]

    cust_col = match_column(headers, CUSTOMER_KEYWORDS)
    addr_col = match_column(headers, ADDRESS_KEYWORDS)
    dist_col = match_column(headers, DISTRICT_KEYWORDS)
    city_col = match_column(headers, CITY_KEYWORDS)
    phone_col = match_column(headers, PHONE_KEYWORDS)
    barcode_col = match_column(headers, BARCODE_KEYWORDS)
    name_col = match_column(headers, NAME_KEYWORDS)
    sku_col = match_column(headers, SKU_KEYWORDS)
    qty_col = match_column(headers, QTY_KEYWORDS)

    if barcode_col is None and sku_col is not None:
        barcode_col = sku_col

    mapping = {
        'cust_idx': cust_col,
        'addr_idx': addr_col,
        'dist_idx': dist_col,
        'city_idx': city_col,
        'phone_idx': phone_col,
        'barcode_idx': barcode_col,
        'name_idx': name_col,
        'sku_idx': sku_col,
        'qty_idx': qty_col,
        'cust_name': headers[cust_col] if cust_col is not None else None,
        'addr_name': headers[addr_col] if addr_col is not None else None,
        'dist_name': headers[dist_col] if dist_col is not None else None,
        'city_name': headers[city_col] if city_col is not None else None,
        'phone_name': headers[phone_col] if phone_col is not None else None,
        'barcode_name': headers[barcode_col] if barcode_col is not None else None,
        'name_name': headers[name_col] if name_col is not None else None,
        'sku_name': headers[sku_col] if sku_col is not None else None,
        'qty_name': headers[qty_col] if qty_col is not None else None
    }

    return headers, data_rows, mapping


def extract_items(headers, data_rows, mapping):
    c_idx = mapping.get('cust_idx')
    a_idx = mapping.get('addr_idx')
    d_idx = mapping.get('dist_idx')
    ct_idx = mapping.get('city_idx')
    p_idx = mapping.get('phone_idx')
    b_idx = mapping.get('barcode_idx')
    n_idx = mapping.get('name_idx')
    s_idx = mapping.get('sku_idx')
    q_idx = mapping.get('qty_idx')

    items = []
    for r_idx, row in enumerate(data_rows):
        def get_val(idx):
            if idx is not None and 0 <= idx < len(row):
                return row[idx]
            return None

        cust = clean_text(get_val(c_idx))
        addr = clean_text(get_val(a_idx))
        dist = clean_text(get_val(d_idx))
        city = clean_text(get_val(ct_idx))
        phone = clean_text(get_val(p_idx))
        barcode = clean_barcode(get_val(b_idx))
        name = clean_text(get_val(n_idx))
        sku = clean_text(get_val(s_idx))
        qty = clean_quantity(get_val(q_idx))

        if not barcode and not name and not cust:
            continue
        if not barcode and sku:
            barcode = sku

        # İlçe içindeki parantezi temizle (Örn: "MERKEZ   (Isparta)" -> dist_clean="MERKEZ", dist_paren="Isparta")
        dist_paren = ""
        match_p = re.search(r'\((.*?)\)', dist)
        if match_p:
            dist_paren = match_p.group(1).strip()
        dist_clean = re.sub(r'\(.*?\)', '', dist).strip()

        addr_line = addr
        if dist_clean and dist_clean.lower() not in addr_line.lower() and len(addr_line) < 55:
            addr_line = f"{addr_line} {dist_clean}".strip()

        city_line = ""
        if dist_paren and city:
            city_line = f"({dist_paren}) {city}".strip()
        elif dist_clean and city:
            city_line = f"({dist_clean}) {city}".strip()
        elif city:
            city_line = city.strip()
        elif dist_clean:
            city_line = dist_clean.strip()

        items.append({
            'id': r_idx + 1,
            'customer': cust,
            'address': addr_line,
            'district': dist,
            'city': city,
            'city_line': city_line,
            'phone': phone,
            'barcode': barcode,
            'name': name or 'İsimsiz Ürün',
            'sku': sku,
            'quantity': qty,
            'selected': True
        })

    return items


def generate_labels_pdf(items, options=None):
    """
    ReportLab ile kullanıcının tam istediği şablonda 100mm x 100mm (10cm x 10cm)
    vektörel kargo & etiket PDF'i üretir.
    Üst: Müşteri Adı, Adres, (İlçe) İl, Telefon
    Orta: Barkod (Code128) ve Barkod Numarası
    Alt (Kırmızı işaretli yer): Ürün Adı
    """
    init_fonts()
    if options is None:
        options = {}

    print_count_mode = options.get('print_count_mode', 'quantity')
    barcode_height_mm = int(options.get('barcode_height', 24))
    font_size_pt = int(options.get('font_size', 12))

    label_size = (100 * mm, 100 * mm)
    margin = 5 * mm

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        pdf_buffer,
        pagesize=label_size,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=5 * mm,
        bottomMargin=5 * mm
    )

    registered = pdfmetrics.getRegisteredFontNames()
    bold_font = 'LabelSegoe-Bold' if 'LabelSegoe-Bold' in registered else ('LabelFont-Bold' if 'LabelFont-Bold' in registered else 'Helvetica-Bold')
    regular_font = 'LabelSegoe' if 'LabelSegoe' in registered else ('LabelFont' if 'LabelFont' in registered else 'Helvetica')

    # Kullanıcı görseliyle birebir font ve hizalama stilleri
    cust_style = ParagraphStyle(
        'CustStyle',
        fontName=bold_font,
        fontSize=13,
        leading=16,
        alignment=1, # Center
        textColor=HexColor('#000000')
    )

    addr_style = ParagraphStyle(
        'AddrStyle',
        fontName=regular_font,
        fontSize=10,
        leading=13,
        alignment=1,
        textColor=HexColor('#111827')
    )

    phone_style = ParagraphStyle(
        'PhoneStyle',
        fontName=regular_font,
        fontSize=10.5,
        leading=14,
        alignment=1,
        textColor=HexColor('#111827')
    )

    # Kırmızı çizilen yer: Ürün Adı Stili
    prod_style = ParagraphStyle(
        'ProdStyle',
        fontName=bold_font,
        fontSize=font_size_pt,
        leading=font_size_pt + 3,
        alignment=1,
        textColor=HexColor('#000000')
    )

    story = []
    printable_items = []
    for item in items:
        if not item.get('selected', True):
            continue
        count = item.get('quantity', 1) if print_count_mode == 'quantity' else 1
        for i in range(count):
            printable_items.append({
                'customer': item.get('customer', ''),
                'address': item.get('address', ''),
                'city_line': item.get('city_line', ''),
                'phone': item.get('phone', ''),
                'barcode': item.get('barcode', ''),
                'name': item.get('name', ''),
                'sku': item.get('sku', ''),
                'qty_current': i + 1,
                'qty_total': count,
                'original_qty': item.get('quantity', 1)
            })

    total_labels = len(printable_items)
    if total_labels == 0:
        story.append(Paragraph("Yazdırılacak etiket bulunamadı.", cust_style))
        doc.build(story)
        pdf_buffer.seek(0)
        return pdf_buffer.getvalue()

    # Sipariş Numarası (Barkod) tekrar sayısı hesabı:
    # SADECE aynı sipariş numarasına sahip etiket sayısı > 1 ise sağ tarafa 1/2, 2/2 yazılır.
    order_totals = {}
    for p in printable_items:
        b = str(p.get('barcode', '')).strip()
        if b:
            order_totals[b] = order_totals.get(b, 0) + 1

    order_counters = {}
    for p in printable_items:
        b = str(p.get('barcode', '')).strip()
        tot = order_totals.get(b, 0)
        if tot > 1:
            cnt = order_counters.get(b, 0) + 1
            order_counters[b] = cnt
            p['package_ratio'] = f"{cnt}/{tot}"
        else:
            p['package_ratio'] = ""

    for idx, p_item in enumerate(printable_items):
        label_elements = []

        # 1. Müşteri Adı Soyadı
        if p_item['customer']:
            label_elements.append(Paragraph(p_item['customer'], cust_style))
            label_elements.append(Spacer(1, 1.5 * mm))

        # 2. Adres
        if p_item['address']:
            label_elements.append(Paragraph(p_item['address'], addr_style))
            label_elements.append(Spacer(1, 1 * mm))

        # 3. (İlçe) İl
        if p_item['city_line']:
            label_elements.append(Paragraph(p_item['city_line'], addr_style))
            label_elements.append(Spacer(1, 1 * mm))

        # 4. Telefon
        if p_item['phone']:
            label_elements.append(Paragraph(p_item['phone'], phone_style))
            label_elements.append(Spacer(1, 2 * mm))
        elif not p_item['customer']:
            label_elements.append(Spacer(1, 4 * mm))

        # 5. Barkod (Code128) ve Sağ Tarafta 1/2, 2/2 (Sadece aynı siparişlerde)
        barcode_val = p_item['barcode']
        package_ratio = p_item.get('package_ratio', '')

        if barcode_val:
            try:
                val_len = len(barcode_val)
                bar_width = 0.42 * mm
                if val_len > 16:
                    bar_width = 0.30 * mm
                elif val_len > 12:
                    bar_width = 0.35 * mm

                bc = createBarcodeDrawing(
                    'Code128',
                    value=barcode_val,
                    barHeight=barcode_height_mm * mm,
                    barWidth=bar_width,
                    humanReadable=True,
                    fontSize=9.5,
                    fontName=regular_font
                )

                if package_ratio:
                    ratio_style = ParagraphStyle(
                        'RatioStyle',
                        fontName=bold_font,
                        fontSize=14,
                        leading=18,
                        alignment=2, # Sağa dayalı
                        textColor=HexColor('#000000')
                    )
                    ratio_p = Paragraph(f"<b>{package_ratio}</b>", ratio_style)
                    t_bc = Table([[Paragraph('', cust_style), bc, ratio_p]], colWidths=[14 * mm, 62 * mm, 14 * mm])
                    t_bc.setStyle(TableStyle([
                        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                        ('LEFTPADDING', (0,0), (-1,-1), 0),
                        ('RIGHTPADDING', (0,0), (-1,-1), 0),
                        ('TOPPADDING', (0,0), (-1,-1), 0),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 0)
                    ]))
                    label_elements.append(t_bc)
                else:
                    bc.hAlign = 'CENTER'
                    label_elements.append(bc)
            except Exception as e:
                err_p = Paragraph(f"Barkod: {barcode_val}", cust_style)
                label_elements.append(err_p)
        else:
            label_elements.append(Paragraph("[ BARKOD YOK ]", cust_style))

        label_elements.append(Spacer(1, 3 * mm))

        # 6. ÜRÜN ADI (Kullanıcının kırmızı daire ile işaretlediği yer)
        clean_prod_name = p_item['name']
        current_prod_style = prod_style
        if len(clean_prod_name) > 75:
            current_prod_style = ParagraphStyle(
                'LongProdStyle',
                parent=prod_style,
                fontSize=max(9, font_size_pt - 2.5),
                leading=max(11, font_size_pt)
            )
        elif len(clean_prod_name) > 45:
            current_prod_style = ParagraphStyle(
                'MedProdStyle',
                parent=prod_style,
                fontSize=max(10, font_size_pt - 1.5),
                leading=max(12, font_size_pt + 1)
            )

        label_elements.append(Paragraph(clean_prod_name, current_prod_style))



        story.append(KeepTogether(label_elements))

        if idx < total_labels - 1:
            story.append(PageBreak())

    doc.build(story)
    pdf_buffer.seek(0)
    return pdf_buffer.getvalue()


def create_sample_excel():
    """Kullanıcının gönderdiği örnek formata tam uygun şablon oluşturur."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kargo Etiket Listesi"

    headers = ["Müşteri Adı", "Adres", "İlçe", "İl", "Telefon", "Barkod", "Ürün Adı", "Stok Kodu", "Adet"]
    ws.append(headers)

    samples = [
        ("mehmet ali tezcan", "çelebiler mahalesi 1409 sokak no:2 MERKEZ", "Isparta", "Isparta", "0(532) 231 78 65", "1092488", "Tefal FV5715 Easygliss Plus 2400 W Buharlı Ütü Kırmızı/Siyah", "2100118363", 1),
        ("Aygül SOYTÜRK", "Kızıltoprak mah 960 sokak no 20 kat 3 daire 9", "Muratpaşa", "Antalya", "0(555) 123 45 67", "7330037481172640", "Tefal Titanyum 6X Excellence İndüksiyon Tabanlı 30 cm Tava", "4300007578", 2),
        ("Şuheda YAVUZ", "Yavuz Selim Mah Gül Gönül Sokak No: 5/4", "Yeşilyurt", "Malatya", "0(544) 987 65 43", "7330037480541447", "Babyliss C325E Sublime Touch Saç Maşası", "8910481200", 1),
        ("Kemal KAYA", "Cumhuriyet Cad. Menekşe Apt. No: 12", "Kadıköy", "İstanbul", "0(533) 444 55 66", "8690842603831", "Beko BKK 2300 Beyaz Mini Telve Türk Kahve Makinesi", "PH-9252", 1)
    ]

    for row in samples:
        ws.append(list(row))

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 50
    ws.column_dimensions['H'].width = 16
    ws.column_dimensions['I'].width = 8

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    return buf.getvalue()
